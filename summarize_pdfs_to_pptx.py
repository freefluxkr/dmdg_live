import os
import sys
import json
import time
import subprocess

# =====================================================================
# 📦 의존성 라이브러리 자동 검사 및 설치
# =====================================================================
def check_dependencies():
    dependencies = {
        "pypdf": "pypdf",
        "pptx": "python-pptx"
    }
    
    for module_name, pip_name in dependencies.items():
        try:
            __import__(module_name)
        except ImportError:
            print(f"💡 [Dependency] '{pip_name}' 라이브러리가 없습니다. 자동 설치를 시작합니다...")
            try:
                subprocess.run([sys.executable, "-m", "pip", "install", pip_name], check=True)
                print(f"✓ '{pip_name}' 설치 완료.")
            except Exception as e:
                print(f"❌ 라이브러리 설치 실패: {e}")
                print(f"👉 수동 설치 명령어: pip install {pip_name}")
                sys.exit(1)

check_dependencies()

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# =====================================================================
# 🎨 PPT 디자인 테마 설정 (기존 테마 계승)
# =====================================================================
FONT_TITLE = "Malgun Gothic"
FONT_BODY = "Malgun Gothic"
FONT_CODE = "Consolas"

COLOR_NAVY = RGBColor(30, 65, 90)         # 주 색상 (헤더, 타이틀 배경)
COLOR_YELLOW = RGBColor(245, 195, 65)     # 포인트 색상 (강조)
COLOR_WHITE = RGBColor(255, 255, 255)     # 흰색
COLOR_CHARCOAL = RGBColor(44, 62, 80)     # 본문 텍스트 색상
COLOR_CODE_BG = RGBColor(40, 44, 52)      # 코드 블록 어두운 배경 (One Dark 테마)
COLOR_LIGHT_GRAY = RGBColor(245, 246, 247) # 슬라이드 기본 연회색 배경

# =====================================================================
# 🎨 PPTX 빌더 함수
# =====================================================================
def build_pptx(ppt_data, output_path):
    prs = Presentation()
    
    # 16:9 와이드스크린 비율 설정
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6] # 완전 빈 슬라이드 레이아웃
    
    ch_num = ppt_data.get("chapter_num", "00")
    ch_title = ppt_data.get("chapter_title", "파이썬 강좌")
    
    # -----------------------------------------------------------------
    # 🎨 [표지 슬라이드] 생성
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(blank_layout)
    
    # 어두운 네이비 배경 생성
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_NAVY
    bg_shape.line.fill.background()
    
    # 노란색 포인트 장식 선
    line_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(2.2), Inches(10.333), Inches(0.1)
    )
    line_shape.fill.solid()
    line_shape.fill.fore_color.rgb = COLOR_YELLOW
    line_shape.line.fill.background()
    
    # 대제목
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10.333), Inches(2.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"Chapter {ch_num}. {ch_title}"
    p.font.name = FONT_TITLE
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.LEFT
    
    # 소제목
    subtitle_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.5), Inches(10.333), Inches(1.0))
    tf_sub = subtitle_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "파이썬 핵심 요약 및 실전 예제 가이드"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = COLOR_YELLOW
    p_sub.alignment = PP_ALIGN.LEFT
    
    # 작성자 정보
    author_box = slide.shapes.add_textbox(Inches(1.5), Inches(5.8), Inches(10.333), Inches(0.8))
    tf_auth = author_box.text_frame
    p_auth = tf_auth.paragraphs[0]
    p_auth.text = "AI 시대의 핵심 프로그래밍 학습 요약 가이드"
    p_auth.font.name = FONT_BODY
    p_auth.font.size = Pt(14)
    p_auth.font.color.rgb = RGBColor(170, 190, 205)
    p_auth.alignment = PP_ALIGN.LEFT
    
    # -----------------------------------------------------------------
    # 🎨 [본문 슬라이드] 생성 루프
    # -----------------------------------------------------------------
    for slide_data in ppt_data.get("slides", []):
        content_slide = prs.slides.add_slide(blank_layout)
        
        # 연한 회색 배경
        c_bg = content_slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        c_bg.fill.solid()
        c_bg.fill.fore_color.rgb = COLOR_LIGHT_GRAY
        c_bg.line.fill.background()
        
        # 1. 상단 타이틀 텍스트 박스
        t_box = content_slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.0))
        t_tf = t_box.text_frame
        t_tf.word_wrap = True
        t_p = t_tf.paragraphs[0]
        t_p.text = slide_data.get("section", "요약 내용")
        t_p.font.name = FONT_TITLE
        t_p.font.size = Pt(24)
        t_p.font.bold = True
        t_p.font.color.rgb = COLOR_NAVY
        
        # 장 카테고리 표시용 서브 텍스트
        sub_lbl_box = content_slide.shapes.add_textbox(Inches(0.8), Inches(0.15), Inches(10.0), Inches(0.4))
        sub_lbl_tf = sub_lbl_box.text_frame
        sub_lbl_p = sub_lbl_tf.paragraphs[0]
        sub_lbl_p.text = f"Chapter {ch_num}. {ch_title}"
        sub_lbl_p.font.name = FONT_BODY
        sub_lbl_p.font.size = Pt(11)
        sub_lbl_p.font.bold = True
        sub_lbl_p.font.color.rgb = RGBColor(140, 150, 160)
        
        # 상단 경계 장식 라인
        dec_line = content_slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.733), Inches(0.04)
        )
        dec_line.fill.solid()
        dec_line.fill.fore_color.rgb = COLOR_NAVY
        dec_line.line.fill.background()
        
        # 코드의 존재 여부에 따라 레이아웃 동적 대응
        code_text = slide_data.get("code", "").strip()
        has_code = len(code_text) > 0
        
        # 2. 핵심 요약 텍스트 영역 (Bullets)
        if has_code:
            text_box = content_slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.5), Inches(5.2))
        else:
            # 코드가 없으면 가로 너비를 꽉 채움
            text_box = content_slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2))
            
        text_tf = text_box.text_frame
        text_tf.word_wrap = True
        
        bullets = slide_data.get("bullets", [])
        for b_idx, bullet in enumerate(bullets):
            p_bullet = text_tf.paragraphs[0] if b_idx == 0 else text_tf.add_paragraph()
            p_bullet.text = bullet
            p_bullet.font.name = FONT_BODY
            p_bullet.font.size = Pt(14)
            p_bullet.font.color.rgb = COLOR_CHARCOAL
            p_bullet.space_after = Pt(14) # 줄 간격 조절
            
            # 서브 불릿 들여쓰기 처리
            if bullet.startswith("  -") or bullet.startswith("  "):
                p_bullet.level = 1
                p_bullet.font.size = Pt(13)
                p_bullet.font.color.rgb = RGBColor(80, 95, 110)
            else:
                p_bullet.level = 0
                
        # 3. 우측: 예제 코드 블록 영역 (코드가 있을 때만 렌더링)
        if has_code:
            # 어두운 배경 상자 그리기
            code_bg_shape = content_slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.5), Inches(5.933), Inches(5.2)
            )
            code_bg_shape.fill.solid()
            code_bg_shape.fill.fore_color.rgb = COLOR_CODE_BG
            code_bg_shape.line.fill.background()
            
            # 박스 상단에 "파이썬 실전 예제" 라벨
            label_box = content_slide.shapes.add_textbox(Inches(6.8), Inches(1.6), Inches(3.0), Inches(0.4))
            label_tf = label_box.text_frame
            label_p = label_tf.paragraphs[0]
            label_p.text = "📎 파이썬 실전 예제"
            label_p.font.name = FONT_BODY
            label_p.font.size = Pt(11)
            label_p.font.bold = True
            label_p.font.color.rgb = COLOR_YELLOW
            
            # 코드 텍스트 박스 얹기
            code_box = content_slide.shapes.add_textbox(Inches(6.8), Inches(2.1), Inches(5.533), Inches(4.5))
            code_tf = code_box.text_frame
            code_tf.word_wrap = True
            code_p = code_tf.paragraphs[0]
            code_p.text = code_text
            code_p.font.name = FONT_CODE
            code_p.font.size = Pt(11)
            code_p.font.color.rgb = COLOR_WHITE
            code_p.line_spacing = 1.2
            
    # PPTX 파일 저장
    prs.save(output_path)

# =====================================================================
# 🚀 메인 실행 루프 (완전 로컬 빌드)
# =====================================================================
def main():
    source_dir = "0704보강"
    target_dir = os.path.join(source_dir, "ppt")
    json_path = os.path.join(source_dir, "ppt_data.json")
    
    if not os.path.exists(source_dir):
        print(f"❌ 에러: 소스 디렉토리 '{source_dir}'가 존재하지 않습니다.")
        sys.exit(1)
        
    if not os.path.exists(json_path):
        print(f"❌ 에러: 요약 데이터 파일 '{json_path}'를 찾을 수 없습니다.")
        print("👉 AI 에이전트가 생성해 둔 ppt_data.json 파일이 작업 공간 내에 있는지 확인해 주세요.")
        sys.exit(1)
        
    if not os.path.exists(target_dir):
        os.makedirs(target_dir)
        print(f"📂 대상 디렉토리 생성됨: {target_dir}")
        
    # 요약 데이터 JSON 로드
    try:
        with open(json_path, 'r', encoding='utf-8') as f:
            full_ppt_data = json.load(f)
        print(f"✓ 요약 데이터 파일 로드 성공 (총 {len(full_ppt_data)}개 챕터)")
    except Exception as e:
        print(f"❌ 요약 데이터 로드 실패: {e}")
        sys.exit(1)
        
    # PDF 파일 목록 수집 및 정렬
    pdf_files = [f for f in os.listdir(source_dir) if f.lower().endswith(".pdf")]
    pdf_files.sort()
    
    if not pdf_files:
        print(f"⚠️ 경고: '{source_dir}' 폴더 내에 PDF 파일이 없습니다.")
        sys.exit(0)
        
    print("=" * 60)
    print(f"🚀 총 {len(pdf_files)}개의 PDF를 고품질 PPTX로 로컬 변환 시작")
    print("=" * 60)
    
    success_count = 0
    for idx, filename in enumerate(pdf_files):
        key_name = os.path.splitext(filename)[0]
        pptx_filename = key_name + ".pptx"
        pptx_path = os.path.join(target_dir, pptx_filename)
        
        print(f"\n[{idx+1}/{len(pdf_files)}] 📄 변환 중: {filename}")
        
        # JSON 요약 데이터에서 해당 장 찾기
        chapter_data = full_ppt_data.get(key_name)
        if not chapter_data:
            print(f"   ⚠️ 경고: JSON 데이터 파일 내에 '{key_name}'에 대한 요약본이 없습니다. 건너뜁니다.")
            continue
            
        # PPTX 빌드 및 저장
        try:
            print(f"   🎨 PPTX 디자인 및 빌드 중...")
            build_pptx(chapter_data, pptx_path)
            print(f"   🎉 PPTX 저장 완료: {pptx_path}")
            success_count += 1
        except Exception as e:
            print(f"   ❌ PPTX 생성 및 저장 실패: {e}")
            continue
            
    print("\n" + "=" * 60)
    print(f"🎉 작업 완료! 총 {success_count}/{len(pdf_files)}개의 PPTX 생성을 마무리했습니다.")
    print(f"📂 결과 폴더: {os.path.abspath(target_dir)}")
    print("=" * 60)

if __name__ == "__main__":
    main()
